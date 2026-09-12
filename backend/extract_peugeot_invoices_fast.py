import os
import re
import glob
from concurrent.futures import ThreadPoolExecutor, as_completed
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import pypdf

BASE_DIR = r"C:\Users\Sanka\OneDrive - George Steuart Optimize Pvt Ltd\Desktop\peugout\New folder"
OUTPUT_EXCEL = r"C:\Users\Sanka\OneDrive - George Steuart Optimize Pvt Ltd\Desktop\peugout\Peugeot_Land_Customer_Invoices_Master.xlsx"

def clean_vehicle_no(raw):
    if not raw:
        return ""
    v = raw.strip().upper()
    v = re.sub(r'PEUGEOT\s*LAND.*', '', v, flags=re.I)
    v = re.sub(r'CUSTOMER\s*DETAILS', '', v, flags=re.I)
    v = re.sub(r'[\r\n\t]', ' ', v)
    v = re.sub(r'\s+', ' ', v).strip()
    v = re.sub(r'\s*-\s*', ' ', v)
    return v

def clean_number(raw):
    if not raw:
        return 0.0
    s = str(raw).replace('LKR', '').replace('Rs', '').replace('rs', '').replace('=', '').replace(' ', '').replace(',', '').strip()
    try:
        return float(s)
    except:
        return 0.0

def clean_mileage(raw):
    if not raw:
        return None
    s = str(raw).upper().replace('KM', '').replace('Km', '').replace('km', '').replace(' ', '').replace(',', '').strip()
    s = re.sub(r'[^0-9]', '', s)
    try:
        val = int(s)
        return val if val > 0 else None
    except:
        return None

def parse_file(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    record = {
        "file_name": os.path.basename(file_path),
        "folder_name": os.path.basename(os.path.dirname(file_path)),
        "file_path": file_path,
        "doc_type": "INVOICE",
        "invoice_no": "",
        "invoice_date": "",
        "vehicle_no": "",
        "make_model": "",
        "mileage": None,
        "chassis_no": "",
        "customer_name": "",
        "customer_phone": "",
        "services_summary": "",
        "parts_summary": "",
        "line_items": [],
        "total_amount": 0.0
    }

    full_text = ""
    if ext == '.pdf':
        try:
            reader = pypdf.PdfReader(file_path)
            full_text = "\n".join([page.extract_text() or "" for page in reader.pages])
        except Exception:
            return None
    elif ext in ['.docx', '.doc']:
        try:
            import docx
            doc = docx.Document(file_path)
            full_text = "\n".join([p.text for p in doc.paragraphs])
            for table in doc.tables:
                for row in table.rows:
                    full_text += "\n" + " ".join([cell.text for cell in row.cells])
        except:
            return None
    elif ext in ['.xlsx', '.xls']:
        fn_veh_m = re.search(r'([A-Z]{2,3}\s*[\-]?\s*[0-9]{3,4})', os.path.basename(file_path), re.I)
        record["vehicle_no"] = clean_vehicle_no(fn_veh_m.group(1)) if fn_veh_m else "UNKNOWN / GENERAL"
        record["invoice_no"] = os.path.splitext(os.path.basename(file_path))[0]
        return record

    if not full_text:
        return None

    # Check Document Type
    if "ESTIMATE" in full_text.upper() or "Estimate No" in full_text:
        record["doc_type"] = "ESTIMATE"

    # Invoice / Estimate Number
    inv_m = re.search(r'(?:Invoice\s*No|Invoice\s*#|Estimate\s*No|Estimate\s*#)\s*[:#]?\s*([A-Za-z0-9\-_]+)', full_text, re.I)
    if inv_m and inv_m.group(1).upper() not in ['DATE', 'TECHNICIAN', 'MILEAGE', 'DESCRIPTION']:
        record["invoice_no"] = inv_m.group(1).strip()
    else:
        fn_m = re.search(r'([0-9]{3,5})', os.path.basename(file_path))
        if fn_m:
            record["invoice_no"] = fn_m.group(1)
        else:
            record["invoice_no"] = os.path.splitext(os.path.basename(file_path))[0]

    # Invoice Date
    date_m = re.search(r'(?:Invoice\s*Date|Estimate\s*Date|Date)\s*[:#]?\s*([0-9]{1,2}[/-][0-9]{1,2}[/-][0-9]{2,4})', full_text, re.I)
    if date_m:
        record["invoice_date"] = date_m.group(1).strip()

    # Vehicle Number
    veh_m = re.search(r'Vehicle\s*(?:No|Number)?\s*[:#]?\s*([A-Za-z0-9\s\-]+?)(?:Mile|Chassis|Date|\n|$)', full_text, re.I)
    if veh_m:
        record["vehicle_no"] = clean_vehicle_no(veh_m.group(1))
    
    if not record["vehicle_no"] or len(record["vehicle_no"]) < 3 or record["vehicle_no"] in ['PEUGEOT', 'DATE']:
        cust_m = re.search(r'Customer\s*Details\s*\n\s*([A-Za-z0-9\s\-]+)', full_text, re.I)
        if cust_m:
            record["vehicle_no"] = clean_vehicle_no(cust_m.group(1))

    if not record["vehicle_no"] or len(record["vehicle_no"]) < 3:
        fn_veh_m = re.search(r'([A-Z]{2,3}\s*[\-]?\s*[0-9]{3,4})', os.path.basename(file_path), re.I)
        if fn_veh_m:
            record["vehicle_no"] = clean_vehicle_no(fn_veh_m.group(1))
        else:
            record["vehicle_no"] = "UNKNOWN / GENERAL"

    # Make / Model
    model_m = re.search(r'Make\s*/?\s*Model\s*[:#]?\s*([A-Za-z0-9\s\-]+?)(?:Chassis|Vehicle|Mile|Date|\n|$)', full_text, re.I)
    if model_m:
        m_str = model_m.group(1).strip()
        if m_str and len(m_str) < 35 and not m_str.upper().startswith('CHASSIS'):
            record["make_model"] = f"Peugeot {m_str}" if "PEUGEOT" not in m_str.upper() else m_str

    # Mileage
    mile_m = re.search(r'Mile\s*age\s*[:;#]?\s*([0-9\s,]+)\s*(?:KM|Km)?', full_text, re.I)
    if mile_m:
        record["mileage"] = clean_mileage(mile_m.group(1))

    # Chassis / VIN
    chassis_m = re.search(r'Chassis\s*(?:no|No)?\s*[:#]?\s*([A-Za-z0-9\s\-]+?)(?:Vehicle|Mile|Date|\n|$)', full_text, re.I)
    if chassis_m:
        c_str = chassis_m.group(1).strip()
        if len(c_str) >= 5:
            record["chassis_no"] = c_str

    # Total Amount
    tot_matches = re.findall(r'(?:Total|LKR|Rs\.?)\s*[:#]?\s*(?:LKR|Rs\.?)?\s*([0-9]{1,3}(?:,[0-9]{2,3})*(?:\.[0-9]{2})?)', full_text, re.I)
    if tot_matches:
        cand_vals = [clean_number(m) for m in tot_matches if clean_number(m) > 100]
        if cand_vals:
            record["total_amount"] = cand_vals[-1]

    # Extract Line descriptions
    lines = full_text.split('\n')
    labour_items = []
    part_items = []
    
    for l in lines:
        l_str = l.strip()
        if len(l_str) < 4 or any(skip in l_str.upper() for skip in ['PEUGEOT LAND', 'GORAKAPOLA', 'PANADURA', 'HOTLINE', 'EMAIL', 'THANK YOU', 'CHEQUE', 'COMPUTER GENERATED', 'SIGNATURE']):
            continue
        
        if any(w in l_str.upper() for w in ['LABOUR', 'CHARGE', 'REPAIR', 'REMOVE', 'SERVICE', 'SCANNER', 'OVERHAUL', 'TEST', 'FIT', 'CLEAN']):
            labour_items.append(l_str)
        elif any(w in l_str.upper() for w in ['OIL', 'FILTER', 'BELT', 'BUSH', 'PAD', 'PLUG', 'PUMP', 'SENSOR', 'VALVE', 'COOLANT', 'SEAL', 'HOSE', 'BATTERY', 'RACK', 'BEARING', 'SHOCK', 'ARM', 'GASKET']):
            part_items.append(l_str)

    record["services_summary"] = " | ".join(labour_items[:4])
    record["parts_summary"] = " | ".join(part_items[:5])

    return record

def main():
    print("Collecting files...")
    all_files = []
    for root, dirs, files in os.walk(BASE_DIR):
        for f in files:
            if not f.startswith('~$') and not f.endswith('.lnk') and not f.endswith('.tmp'):
                all_files.append(os.path.join(root, f))

    print(f"Total files to parse: {len(all_files)}")

    invoices_data = []
    with ThreadPoolExecutor(max_workers=16) as executor:
        futures = {executor.submit(parse_file, f): f for f in all_files}
        for future in as_completed(futures):
            res = future.result()
            if res:
                invoices_data.append(res)

    print(f"Successfully extracted {len(invoices_data)} invoice documents!")

    # ── AGGREGATE BY VEHICLE NUMBER (PRIMARY KEY) ──────────────────────────
    vehicles_dict = {}
    for inv in invoices_data:
        v_num = inv["vehicle_no"]
        if v_num not in vehicles_dict:
            vehicles_dict[v_num] = {
                "Vehicle_Number": v_num,
                "Make_Model": inv["make_model"] or "Peugeot",
                "Customer_Name": inv["customer_name"] or "",
                "Customer_Phone": inv["customer_phone"] or "",
                "Total_Invoices_Count": 0,
                "Total_Lifetime_Spend_LKR": 0.0,
                "Earliest_Visit_Date": inv["invoice_date"] or "",
                "Latest_Visit_Date": inv["invoice_date"] or "",
                "Latest_Mileage_KM": inv["mileage"] or None,
                "Chassis_VIN": inv["chassis_no"] or "",
                "Services_History": []
            }

        v = vehicles_dict[v_num]
        v["Total_Invoices_Count"] += 1
        v["Total_Lifetime_Spend_LKR"] += inv["total_amount"]
        if not v["Make_Model"] and inv["make_model"]:
            v["Make_Model"] = inv["make_model"]
        if not v["Chassis_VIN"] and inv["chassis_no"]:
            v["Chassis_VIN"] = inv["chassis_no"]
        if inv["mileage"]:
            if not v["Latest_Mileage_KM"] or inv["mileage"] > v["Latest_Mileage_KM"]:
                v["Latest_Mileage_KM"] = inv["mileage"]
        if inv["invoice_date"]:
            v["Latest_Visit_Date"] = inv["invoice_date"]
        
        all_desc = ([inv["services_summary"]] if inv["services_summary"] else []) + ([inv["parts_summary"]] if inv["parts_summary"] else [])
        for d in all_desc:
            if d and d not in v["Services_History"]:
                v["Services_History"].append(d)

    vehicles_rows = []
    for v_num, d in vehicles_dict.items():
        vehicles_rows.append({
            "Vehicle_Number": d["Vehicle_Number"],
            "Make_Model": d["Make_Model"],
            "Customer_Name": d["Customer_Name"],
            "Customer_Phone": d["Customer_Phone"],
            "Total_Invoices_Count": d["Total_Invoices_Count"],
            "Total_Lifetime_Spend_LKR": round(d["Total_Lifetime_Spend_LKR"], 2),
            "Earliest_Visit_Date": d["Earliest_Visit_Date"],
            "Latest_Visit_Date": d["Latest_Visit_Date"],
            "Latest_Mileage_KM": d["Latest_Mileage_KM"],
            "Chassis_VIN": d["Chassis_VIN"],
            "Key_Services_Done": " | ".join(d["Services_History"][:3])
        })

    vehicles_rows.sort(key=lambda x: x["Total_Invoices_Count"], reverse=True)

    invoices_rows = []
    for inv in invoices_data:
        invoices_rows.append({
            "Invoice_Number": inv["invoice_no"],
            "Document_Type": inv["doc_type"],
            "Invoice_Date": inv["invoice_date"],
            "Vehicle_Number": inv["vehicle_no"],
            "Make_Model": inv["make_model"],
            "Mileage_KM": inv["mileage"],
            "Replaced_Parts_Summary": inv["parts_summary"],
            "Labour_Services_Summary": inv["services_summary"],
            "Total_Amount_LKR": round(inv["total_amount"], 2),
            "Month_Folder": inv["folder_name"],
            "File_Name": inv["file_name"]
        })

    invoices_rows.sort(key=lambda x: str(x["Invoice_Date"]), reverse=True)

    df_vehicles = pd.DataFrame(vehicles_rows)
    df_invoices = pd.DataFrame(invoices_rows)

    print(f"Writing Excel with {len(df_vehicles)} unique vehicles and {len(df_invoices)} invoices...")

    with pd.ExcelWriter(OUTPUT_EXCEL, engine='openpyxl') as writer:
        df_vehicles.to_excel(writer, sheet_name='Customer_Vehicles_Master', index=False)
        df_invoices.to_excel(writer, sheet_name='All_Invoices_History', index=False)

    # Style Excel
    wb = openpyxl.load_workbook(OUTPUT_EXCEL)
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    for ws in wb.worksheets:
        ws.views.sheetView[0].showGridLines = True
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 50)

        ws.row_dimensions[1].height = 28

    wb.save(OUTPUT_EXCEL)
    print(f"SUCCESS! Master Excel file saved to: {OUTPUT_EXCEL}")

if __name__ == "__main__":
    main()
