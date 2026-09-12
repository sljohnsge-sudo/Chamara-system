import os
import re
import json
import openpyxl
from datetime import datetime
from sqlalchemy import create_engine, text
from sqlalchemy.orm import sessionmaker

DATABASE_URL = "mysql+pymysql://root:@localhost:3306/peugeot_land"
EXCEL_PATH = r"C:\Users\Sanka\OneDrive - George Steuart Optimize Pvt Ltd\Desktop\peugout\New format\New folder\New folder\Peugeot_Invoices_Formatted_Updated.xlsx"

engine = create_engine(DATABASE_URL)
SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)

def clean_num(val):
    if val is None or val == '':
        return 0.0
    try:
        s = str(val).replace('LKR', '').replace('Rs', '').replace(',', '').replace(' ', '').strip()
        return float(s)
    except:
        return 0.0

def clean_qty(val):
    if val is None or val == '':
        return 1
    try:
        s = str(val).strip()
        nums = re.findall(r'[0-9]+(?:\.[0-9]+)?', s)
        if nums:
            return max(1, int(float(nums[0])))
        return 1
    except:
        return 1

def is_labour_item(desc):
    if not desc:
        return False
    d = desc.upper()
    keywords = [
        'LABOUR', 'LABOURE', 'LABOR', 'CHARGE', 'REMOVE & REFIX', 
        'REMOVE & RE-FIX', 'SCANNER', 'LATHE WORK', 'COOLENT FLUSH', 
        'HEAD FACE', 'REPAIRE ONE', 'PIC ME', 'DISCOUNT'
    ]
    return any(k in d for k in keywords)

def main():
    print(f"Loading Excel file from: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb['Sheet1']
    
    # ── 1. GROUP ROWS BY INVOICE ──────────────────────────────────────────
    invoices = {}
    
    for r in range(4, ws.max_row): # exclude bottom summary row
        inv_no = str(ws.cell(r, 2).value or '').strip()
        inv_date = str(ws.cell(r, 3).value or '').strip()
        model = str(ws.cell(r, 4).value or '').strip()
        veh_no = str(ws.cell(r, 5).value or '').strip().upper()
        chassis = str(ws.cell(r, 6).value or '').strip()
        mileage = str(ws.cell(r, 7).value or '').strip()
        desc = str(ws.cell(r, 8).value or '').strip()
        qt_raw = ws.cell(r, 9).value
        unit_raw = ws.cell(r, 10).value
        tot_raw = ws.cell(r, 11).value
        inv_tot_raw = ws.cell(r, 12).value
        
        if not inv_no or not veh_no or desc.upper().startswith('TOTAL WORKSHOP'):
            continue
            
        key = (inv_no, veh_no)
        if key not in invoices:
            invoices[key] = {
                'invoice_no': inv_no,
                'invoice_date': inv_date if inv_date else '2025-01-01',
                'make_model': model if model else 'Peugeot',
                'vehicle_no': veh_no,
                'chassis_no': chassis if chassis and len(chassis) > 4 else None,
                'mileage': mileage if mileage else '0',
                'grand_total': clean_num(inv_tot_raw),
                'items': []
            }
            
        line_tot = clean_num(tot_raw)
        unit_price = clean_num(unit_raw)
        qt = clean_qty(qt_raw)
        
        if line_tot == 0.0 and unit_price > 0:
            line_tot = unit_price * qt
        elif unit_price == 0.0 and line_tot > 0:
            unit_price = round(line_tot / qt, 2)
            
        invoices[key]['items'].append({
            'desc': desc,
            'qt': qt,
            'unit': unit_price,
            'total_rs': line_tot,
            'is_labour': is_labour_item(desc)
        })

    print(f"Loaded {len(invoices)} unique invoices from Excel.")

    # ── 2. FLASH / FLUSH ALL DATA IN DATABASE ──────────────────────────────
    print("\nFlushing existing data from peugeot_land database...")
    with engine.begin() as conn:
        conn.execute(text("SET FOREIGN_KEY_CHECKS = 0;"))
        conn.execute(text("TRUNCATE TABLE job_inventory_issues;"))
        conn.execute(text("TRUNCATE TABLE quotations;"))
        conn.execute(text("TRUNCATE TABLE job_cards;"))
        conn.execute(text("TRUNCATE TABLE vehicles;"))
        conn.execute(text("TRUNCATE TABLE customers;"))
        conn.execute(text("TRUNCATE TABLE inventory_items;"))
        conn.execute(text("SET FOREIGN_KEY_CHECKS = 1;"))
    print("Database flushed successfully!")

    # ── 3. INGEST DATA INTO MYSQL ──────────────────────────────────────────
    session = SessionLocal()
    
    customers_map = {} # veh_no -> customer_id
    vehicles_map = {}  # veh_no -> vehicle_id
    inventory_map = {} # item_name_lower -> item_id
    
    customer_count = 0
    vehicle_count = 0
    job_count = 0
    item_count = 0
    issue_count = 0
    total_sales_ingested = 0.0

    try:
        for (inv_no, veh_no), inv_data in invoices.items():
            # 1. Customer
            if veh_no not in customers_map:
                phone_num = f"077{1000000 + customer_count % 9000000}"
                c_res = session.execute(
                    text("""
                        INSERT INTO customers (name, phone, email, address, created_at)
                        VALUES (:name, :phone, :email, :address, :created_at)
                    """),
                    {
                        "name": f"Customer ({veh_no})",
                        "phone": phone_num,
                        "email": f"{veh_no.replace(' ', '').lower()}@customer.peugeotland.lk",
                        "address": "Colombo, Sri Lanka",
                        "created_at": inv_data['invoice_date']
                    }
                )
                cid = c_res.lastrowid
                customers_map[veh_no] = cid
                customer_count += 1
            else:
                cid = customers_map[veh_no]

            # 2. Vehicle
            if veh_no not in vehicles_map:
                v_res = session.execute(
                    text("""
                        INSERT INTO vehicles (customer_id, vehicle_number, make_model, vin_chassis, created_at)
                        VALUES (:cid, :vnum, :model, :chassis, :created_at)
                    """),
                    {
                        "cid": cid,
                        "vnum": veh_no,
                        "model": inv_data['make_model'],
                        "chassis": inv_data['chassis_no'],
                        "created_at": inv_data['invoice_date']
                    }
                )
                vid = v_res.lastrowid
                vehicles_map[veh_no] = vid
                vehicle_count += 1
            else:
                vid = vehicles_map[veh_no]

            # 3. Calculate Labour & Parts breakdown
            labour_items = [it for it in inv_data['items'] if it['is_labour']]
            parts_items = [it for it in inv_data['items'] if not it['is_labour']]
            
            labour_sum = sum(it['total_rs'] for it in labour_items)
            parts_sum = sum(it['total_rs'] for it in parts_items)
            
            if not parts_items and not labour_items:
                inv_total = inv_data['grand_total']
                labour_sum = inv_total
            else:
                inv_total = parts_sum + labour_sum
                if inv_total == 0.0 and inv_data['grand_total'] > 0:
                    inv_total = inv_data['grand_total']
                    labour_sum = inv_total

            total_sales_ingested += inv_total

            # Job Descriptions
            all_descs = [it['desc'] for it in inv_data['items'] if it['desc']]
            repair_fault_txt = " | ".join(all_descs[:6]) if all_descs else "General Scheduled Maintenance"

            labour_details_json = json.dumps([
                {"description": l['desc'], "amount": l['total_rs']} for l in labour_items
            ]) if labour_items else None

            # 4. Insert Job Card
            job_res = session.execute(
                text("""
                    INSERT INTO job_cards (
                        job_number, vehicle_id, customer_id, mileage, repair_fault, technician_notes, 
                        labour_charge, status, created_at, updated_at, 
                        payment_status, payment_method, payment_received_at, 
                        car_released, car_released_at, labour_details
                    )
                    VALUES (
                        :job_no, :vid, :cid, :mileage, :fault, :notes,
                        :labour, 'Delivered', :created_at, :created_at,
                        'paid', 'Cheque / Cash', :created_at,
                        1, :created_at, :labour_details
                    )
                """),
                {
                    "job_no": f"INV-{inv_no}",
                    "vid": vid,
                    "cid": cid,
                    "mileage": inv_data['mileage'],
                    "fault": repair_fault_txt,
                    "notes": f"Historical Service Invoice #{inv_no}",
                    "labour": labour_sum,
                    "created_at": inv_data['invoice_date'],
                    "labour_details": labour_details_json
                }
            )
            jid = job_res.lastrowid
            job_count += 1

            # 5. Insert Parts & Job Inventory Issues
            for p in parts_items:
                p_desc = p['desc'].strip()
                if not p_desc: continue
                p_key = p_desc.lower()
                
                if p_key not in inventory_map:
                    part_code = f"PART-{1000 + item_count}"
                    p_res = session.execute(
                        text("""
                            INSERT INTO inventory_items (
                                item_code, item_name, applicable_model, 
                                quantity_in_stock, unit_price, cost_price, 
                                unit_of_measure, markup_type, markup_value, created_at
                            )
                            VALUES (
                                :code, :name, :model, 
                                50, :unit_p, :unit_p,
                                'nos', 'percentage', 15.0, :created_at
                            )
                        """),
                        {
                            "code": part_code,
                            "name": p_desc,
                            "model": inv_data['make_model'],
                            "unit_p": p['unit'],
                            "created_at": inv_data['invoice_date']
                        }
                    )
                    part_id = p_res.lastrowid
                    inventory_map[p_key] = part_id
                    item_count += 1
                else:
                    part_id = inventory_map[p_key]

                # Insert Job Inventory Issue
                session.execute(
                    text("""
                        INSERT INTO job_inventory_issues (
                            job_card_id, inventory_item_id, quantity_issued, unit_price, total_price, issued_at
                        )
                        VALUES (
                            :jid, :part_id, :qty, :unit_p, :tot_p, :issued_at
                        )
                    """),
                    {
                        "jid": jid,
                        "part_id": part_id,
                        "qty": p['qt'],
                        "unit_p": p['unit'],
                        "tot_p": p['total_rs'],
                        "issued_at": inv_data['invoice_date']
                    }
                )
                issue_count += 1

        session.commit()
        print("\n=======================================================")
        print("DATABASE RE-INGESTION COMPLETED SUCCESSFULLY!")
        print(f"  Total Customers Created:    {customer_count}")
        print(f"  Total Vehicles Registered:  {vehicle_count}")
        print(f"  Total Job Cards Created:    {job_count}")
        print(f"  Total Cataloged Parts:      {item_count}")
        print(f"  Total Parts Issued to Jobs: {issue_count}")
        print(f"  Total Lifetime Sales:       LKR {total_sales_ingested:,.2f}")
        print("=======================================================\n")

    except Exception as e:
        session.rollback()
        print(f"Error during ingestion: {e}")
        raise e
    finally:
        session.close()

if __name__ == "__main__":
    main()
